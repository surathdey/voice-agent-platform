"""Campaigns & Leads API (A3) — design §8.5 / §11.2 row A3, §15 rules.

A campaign is bound to exactly ONE agent and carries BUSINESS configuration
only: product / purpose / audience, a named CHANNEL (never a phone number),
provider mode, schedule and retry rules. Nothing agent-level (model, voice,
language rule, the single prompt) is settable here — campaigns reference the
agent; the prompt is edited only in the Agent Config Studio (A2).

Leads arrive as CSV — one separate list per campaign, in the format of the
campaign's agent (RPI: callee_name · phone_number · bank_name · state; the
bank allocation format: ACCOUNT_NO · CUSTOMER_NAME · PRINCIPAL · OUTSTANDING
· MOBILE_NO · RegistrationNumber). Upload is two steps: `preview` (headers,
sample rows, suggested mapping) then `upload` (file + confirmed mapping).
The mapped product-identifier column is stored in agent_fields under the
canonical keys product_identifier / product_identifier_type, typed by the
campaign's product (design §15).
"""
import csv
import io
import json
import re

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from .. import providers as prov
from ..db import get_db
from ..deps import get_current_user
from ..models import AgentConfig, CallResult, CallSession, Campaign, Lead, User
from ..schemas import (
    PRODUCTS,
    CampaignIn,
    CampaignOut,
    CampaignStatusIn,
    LeadListOut,
    LeadOut,
    UploadResultOut,
)

router = APIRouter(prefix="/api/campaigns", tags=["campaigns"])

# Status lifecycle (A3): draft → active ⇄ paused → completed (terminal).
TRANSITIONS = {
    "draft": {"active"},
    "active": {"paused", "completed"},
    "paused": {"active", "completed"},
    "completed": set(),
}

CORE_TARGETS = {"name", "phone_number", "bank_name", "state", "product_identifier"}
ALL_TARGETS = CORE_TARGETS | {"agent_field", "ignore"}

# Header aliases used to SUGGEST a mapping in the preview (operator confirms).
ALIASES = {
    "name": {"callee_name", "customer_name", "name", "customer", "full_name"},
    "phone_number": {"phone_number", "mobile_no", "mobile", "phone", "contact_number", "mobile_number"},
    "bank_name": {"bank_name", "bank", "lender_name", "institution"},
    "state": {"state", "state_province", "province", "state/province"},
}
IDENTIFIER_ALIASES = {
    "registration_number": {"registrationnumber", "registration_number", "reg_no", "vehicle_no", "registration_no"},
    "card_last4": {"card_last4", "card_last_4", "last4", "card_number_last4"},
    "account_number": {"account_no", "account_number", "loan_account_no", "account", "loan_no"},
}


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", h.strip().lower()).strip("_")


def _suggest_mapping(headers: list[str], identifier_type: str) -> dict:
    """column name -> target, guessed from header aliases (preview only)."""
    out = {}
    used = set()
    id_aliases = IDENTIFIER_ALIASES.get(identifier_type, set())
    for h in headers:
        n = _norm_header(h)
        target = "agent_field"
        if n in id_aliases:
            target = "product_identifier"
        else:
            for core, aliases in ALIASES.items():
                if n in aliases and core not in used:
                    target = core
                    break
        if target in CORE_TARGETS:
            if target in used:  # two columns claim one core field — keep data instead
                target = "agent_field"
            else:
                used.add(target)
        out[h] = target
    return out


def _read_csv(raw: bytes) -> tuple[list[str], list[dict]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "CSV has no header row")
    headers = [h for h in reader.fieldnames if h is not None]
    rows = [{h: (row.get(h) or "").strip() for h in headers} for row in reader]
    return headers, rows


def _get_scoped(db: Session, user: User, campaign_id: str) -> Campaign:
    c = (
        db.query(Campaign)
        .filter(Campaign.campaign_id == campaign_id, Campaign.tenant_id == user.tenant_id)
        .first()
    )
    if c is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Campaign not found")
    return c


def _agent_scoped(db: Session, user: User, agent_id: str) -> AgentConfig:
    a = (
        db.query(AgentConfig)
        .filter(AgentConfig.agent_id == agent_id, AgentConfig.tenant_id == user.tenant_id)
        .first()
    )
    if a is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "agent_id does not name an agent of this tenant")
    return a


def _out(db: Session, c: Campaign) -> CampaignOut:
    agent = db.get(AgentConfig, c.agent_id)
    meta = PRODUCTS.get(c.product, PRODUCTS["other"])
    # A4: flag campaigns whose mode has no working provider — creating and
    # running them is allowed, but their calls run in stub mode.
    ready, missing = prov.mode_readiness(db, c.tenant_id, c.provider_mode)
    flag = None
    if not ready:
        mode_label = "managed" if c.provider_mode == "managed" else "self-hosted"
        flag = (f"No working provider configured for {mode_label} mode "
                f"(missing: {', '.join(missing)}) — calls will run in stub mode. "
                "Configure it under Settings → Providers.")
    return CampaignOut(
        campaign_id=c.campaign_id, tenant_id=c.tenant_id,
        agent_id=c.agent_id, agent_name=agent.name if agent else "",
        name=c.name, product=c.product, product_label=meta["label"],
        product_detail=c.product_detail, purpose=c.purpose, audience=c.audience,
        channel=c.channel, provider_mode=c.provider_mode, status=c.status,
        settlement_authority_pct=c.settlement_authority_pct,
        identifier_type=meta["identifier_type"], identifier_label=meta["identifier_label"],
        schedule=c.schedule or {}, retry_rules=c.retry_rules or {},
        routing_rules=c.routing_rules or [],
        lead_count=db.query(Lead).filter(Lead.campaign_id == c.campaign_id).count(),
        provider_configured=ready, provider_flag=flag,
        created_at=c.created_at.isoformat() if c.created_at else None,
    )


def _apply(c: Campaign, body: CampaignIn) -> None:
    c.agent_id = body.agent_id
    c.name = body.name
    c.product = body.product
    c.product_detail = body.product_detail
    c.purpose = body.purpose
    c.audience = body.audience
    c.channel = body.channel
    c.provider_mode = body.provider_mode
    c.settlement_authority_pct = body.settlement_authority_pct
    c.schedule = body.schedule.model_dump()
    c.retry_rules = body.retry_rules.model_dump()
    c.routing_rules = [r.model_dump() for r in body.routing_rules]


def _validate(body: CampaignIn) -> None:
    if body.product not in PRODUCTS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"product must be one of {sorted(PRODUCTS)}")
    if body.provider_mode not in ("managed", "self_hosted"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "provider_mode must be 'managed' or 'self_hosted'")
    if body.settlement_authority_pct is not None and not 0 <= body.settlement_authority_pct <= 100:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "settlement_authority_pct must be between 0 and 100")


@router.get("", response_model=list[CampaignOut])
def list_campaigns(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = (
        db.query(Campaign)
        .filter(Campaign.tenant_id == user.tenant_id)
        .order_by(Campaign.created_at)
        .all()
    )
    return [_out(db, c) for c in rows]


@router.post("", response_model=CampaignOut, status_code=status.HTTP_201_CREATED)
def create_campaign(body: CampaignIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    _validate(body)
    _agent_scoped(db, user, body.agent_id)  # bound agent must be this tenant's
    c = Campaign(tenant_id=user.tenant_id, status="draft")
    _apply(c, body)
    db.add(c)
    db.commit()
    db.refresh(c)
    return _out(db, c)


@router.get("/{campaign_id}", response_model=CampaignOut)
def get_campaign(campaign_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return _out(db, _get_scoped(db, user, campaign_id))


@router.put("/{campaign_id}", response_model=CampaignOut)
def update_campaign(campaign_id: str, body: CampaignIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    c = _get_scoped(db, user, campaign_id)
    _validate(body)
    _agent_scoped(db, user, body.agent_id)
    _apply(c, body)
    db.commit()
    db.refresh(c)
    return _out(db, c)


@router.delete("/{campaign_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_campaign(campaign_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    c = _get_scoped(db, user, campaign_id)
    # Delete in foreign-key order: call results -> call sessions -> leads ->
    # campaign. Sessions reference both the campaign and its leads, results
    # reference sessions; Postgres enforces those FKs, so deleting the leads
    # first failed (500) on any campaign that had ever been dispatched.
    lead_ids = [lid for (lid,) in db.query(Lead.lead_id).filter(Lead.campaign_id == c.campaign_id)]
    session_q = db.query(CallSession.session_id).filter(CallSession.campaign_id == c.campaign_id)
    if lead_ids:
        session_q = db.query(CallSession.session_id).filter(
            (CallSession.campaign_id == c.campaign_id) | (CallSession.lead_id.in_(lead_ids))
        )
    session_ids = [sid for (sid,) in session_q]
    if session_ids:
        db.query(CallResult).filter(CallResult.session_id.in_(session_ids)).delete(synchronize_session=False)
        db.query(CallSession).filter(CallSession.session_id.in_(session_ids)).delete(synchronize_session=False)
    db.query(Lead).filter(Lead.campaign_id == c.campaign_id).delete(synchronize_session=False)
    db.delete(c)
    db.commit()
    return None


@router.post("/{campaign_id}/status", response_model=CampaignOut)
def set_status(campaign_id: str, body: CampaignStatusIn, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    c = _get_scoped(db, user, campaign_id)
    allowed = TRANSITIONS.get(c.status, set())
    if body.status not in allowed:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Cannot move a {c.status} campaign to {body.status}; allowed: {sorted(allowed) or 'none'}",
        )
    c.status = body.status
    db.commit()
    db.refresh(c)
    return _out(db, c)


# ── Leads ─────────────────────────────────────────────────────────────────

def _lead_out(l: Lead) -> LeadOut:
    return LeadOut(
        lead_id=l.lead_id, campaign_id=l.campaign_id, phone_number=l.phone_number,
        name=l.name, bank_name=l.bank_name, state=l.state,
        agent_fields=l.agent_fields or {}, attempt_count=l.attempt_count,
        status=l.status,
        next_attempt_at=l.next_attempt_at.isoformat() if l.next_attempt_at else None,
        created_at=l.created_at.isoformat() if l.created_at else None,
    )


@router.get("/{campaign_id}/leads", response_model=LeadListOut)
def list_leads(campaign_id: str, limit: int = 500, offset: int = 0,
               user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    c = _get_scoped(db, user, campaign_id)  # 404 unless this tenant's campaign
    q = db.query(Lead).filter(Lead.campaign_id == c.campaign_id)
    total = q.count()
    rows = q.order_by(Lead.created_at, Lead.lead_id).offset(offset).limit(limit).all()
    return LeadListOut(total=total, leads=[_lead_out(l) for l in rows])


@router.get("/{campaign_id}/export")
def export_campaign(campaign_id: str,
                    user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """The whole campaign as a spreadsheet (xlsx): campaign details,
    one row per lead with its latest result, and one row per call with
    recording link and full transcript. The Calls sheet matches the
    QGBS output columns exactly (ACCOUNT_NO through Recording). The
    operator's download for a finished campaign — the same §8.2 records
    the detail page shows."""
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR,
                            "Spreadsheet export needs the openpyxl package")
    c = _get_scoped(db, user, campaign_id)  # 404 unless this tenant's campaign
    agent = db.get(AgentConfig, c.agent_id)
    leads = (db.query(Lead).filter(Lead.campaign_id == c.campaign_id)
             .order_by(Lead.created_at, Lead.lead_id).all())
    sessions = (db.query(CallSession).filter(CallSession.campaign_id == c.campaign_id)
                .order_by(CallSession.started_at, CallSession.session_id).all())
    results = {}
    if sessions:
        for r in (db.query(CallResult)
                  .filter(CallResult.session_id.in_(
                      [s.session_id for s in sessions])).all()):
            results[r.session_id] = r
    latest = {}
    for s in sessions:
        latest[s.lead_id] = s  # ordered by started_at: last wins
    connected = sum(1 for s in sessions
                    if s.session_id in results and s.state == "COMPLETED")
    disp_counts: dict[str, int] = {}
    total_cost = 0.0
    for s in sessions:
        r = results.get(s.session_id)
        if r is None:
            continue
        disp_counts[r.disposition or "(none)"] = disp_counts.get(r.disposition or "(none)", 0) + 1
        cb = r.cost_breakdown or {}
        if cb.get("mode") != "stub":
            total_cost += float(cb.get("total") or 0.0)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campaign"
    sched = c.schedule or {}
    retry = c.retry_rules or {}
    kv = [
        ("Name", c.name),
        ("Agent", agent.name if agent else ""),
        ("Product", c.product),
        ("Channel", c.channel),
        ("Mode", c.provider_mode),
        ("Status", c.status),
        ("Schedule", f"{sched.get('calling_days', '')} "
                     f"{sched.get('window_start', '')}-{sched.get('window_end', '')}"),
        ("Retries", f"up to {retry.get('max_attempts', '')} per lead, "
                    f"gap {retry.get('retry_gap_days', '')} day(s)"),
        ("Leads", len(leads)),
        ("Attempts", len(sessions)),
        ("Connected", connected),
        ("Total cost (Vapi, USD)", round(total_cost, 4)),
    ]
    for k, v in kv:
        ws.append([k, v])
    for name, n in sorted(disp_counts.items()):
        ws.append([f"Disposition: {name}", n])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=1):
        row[0].font = Font(bold=True)

    ws = wb.create_sheet("Leads")
    ws.append(["Name", "Phone", "Bank", "State", "Identifier", "Other fields",
               "Status", "Next attempt", "Attempts", "Last disposition",
               "Last duration (s)", "Last cost (Vapi, USD)"])
    for l in leads:
        s = latest.get(l.lead_id)
        r = results.get(s.session_id) if s else None
        cb = r.cost_breakdown if r else {}
        others = "; ".join(f"{k}: {v}" for k, v in (l.agent_fields or {}).items()
                           if k not in ("product_identifier", "product_identifier_type"))
        ws.append([l.name or "", l.phone_number or "", l.bank_name or "",
                   l.state or "", (l.agent_fields or {}).get("product_identifier", ""),
                   others, l.status or "",
                   l.next_attempt_at.isoformat() if l.next_attempt_at else "",
                   l.attempt_count or 0, r.disposition if r else "",
                   r.duration_s if r else "",
                   cb.get("total", "") if r else ""])
    for col, w in zip("ABCDEFGHIJKL", [20, 18, 18, 14, 20, 40, 12, 20, 10, 22, 16, 14]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Calls")
    # Calls sheet matches the QGBS output columns EXACTLY (same headers,
    # same order): ACCOUNT_NO, CUSTOMER_NAME, PRINCIPAL, OUTSTANDING,
    # MOBILE_NO, RegistrationNumber, Call ID, Start Time, End Time,
    # Duration, Call Status, Disposition, Category, Summary, Transcript,
    # Recording. Category follows the QGBS mapping derived from their
    # data (ACCEPTED -> Settlement, CALLBACK -> Callback, PRIOR_PAYMENT
    # -> Payment, everything else -> Other). Transcript is the full
    # speaker-tagged text (QGBS links to transcript files; we host no
    # transcript files, so the text itself is the honest equivalent).
    # Recording is the S3 presigned link when archived, else the Vapi
    # URL as reference (playback lives in the call drawer via proxy).
    from datetime import timedelta
    ws.append(["ACCOUNT_NO", "CUSTOMER_NAME", "PRINCIPAL", "OUTSTANDING",
               "MOBILE_NO", "RegistrationNumber", "Call ID", "Start Time",
               "End Time", "Duration", "Call Status", "Disposition",
               "Category", "Summary", "Transcript", "Recording"])
    lead_by_id = {l.lead_id: l for l in leads}
    # S3 links (012) are presigned AT EXPORT TIME (7 days) — never
    # stored, so an exported link is always fresh. Without a bucket or
    # boto3 the column stays blank and the export still succeeds.
    try:
        from .. import s3_recordings as _s3

        _s3_on = bool(_s3.bucket_name())
    except Exception:
        _s3 = None
        _s3_on = False
    _status_map = {"COMPLETED": "Completed", "FAILED": "Failed",
                   "NO_ANSWER": "No Answer", "BUSY": "Busy",
                   "VOICEMAIL": "Voicemail"}
    _category_map = {"SETTLEMENT_OFFER_ACCEPTED": "Settlement",
                     "CALLBACK_REQUESTED": "Callback",
                     "PRIOR_PAYMENT_CLAIMED": "Payment"}

    def _num(v):
        try:
            return float(v) if v is not None and str(v).strip() != "" else ""
        except (TypeError, ValueError):
            return v

    for s in sessions:
        r = results.get(s.session_id)
        turns = s.transcript or []
        text = "\n".join(
            f"{'AGENT' if t.get('speaker') == 'agent' else 'CUSTOMER'}: {t.get('text', '')}"
            for t in turns)
        l = lead_by_id.get(s.lead_id)
        af = (l.agent_fields or {}) if l else {}
        phone = (l.phone_number if l and l.phone_number else s.to_number) or ""
        digits = phone.lstrip("+").strip()
        mobile = int(digits) if digits.isdigit() else phone
        reg = af.get("product_identifier", "") \
            if af.get("product_identifier_type") == "registration_number" else ""
        acct = af.get("account_no") or af.get("account_number") or (
            af.get("product_identifier", "")
            if af.get("product_identifier_type") == "account_number" else "")
        disp = r.disposition if r else ""
        cat = _category_map.get(disp, "Other") if disp else ""
        dur = r.duration_s if r else ""
        end = ""
        if s.started_at and dur != "" and dur is not None:
            try:
                end = (s.started_at + timedelta(seconds=float(dur))).isoformat()
            except (TypeError, ValueError):
                end = ""
        s3_link = ""
        if _s3_on and s.recording_s3_key:
            s3_link = _s3.presigned_playback_url(
                s.recording_s3_bucket, s.recording_s3_key)
        ws.append([str(acct) if acct != "" else "", l.name if l else "",
                   _num(af.get("principal")), _num(af.get("outstanding")),
                   mobile, reg, s.session_id,
                   s.started_at.isoformat() if s.started_at else "", end,
                   dur, _status_map.get(s.state or "", s.state or ""),
                   disp, cat, r.summary if r else "", text,
                   s3_link or s.recording_url or ""])
    for col, w in zip("ABCDEFGHIJKLMNOP",
                      [22, 20, 13, 13, 18, 18, 24, 28, 28, 10, 13, 26, 12, 60, 100, 60]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)
    wb.active = wb.index(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", c.name) or "campaign"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe}_export.xlsx"'})


@router.post("/{campaign_id}/leads/preview")
async def preview_leads(campaign_id: str, file: UploadFile = File(...),
                        user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Parse an uploaded CSV without storing anything: headers, row count,
    sample rows and a SUGGESTED column mapping for the operator to confirm."""
    c = _get_scoped(db, user, campaign_id)
    headers, rows = _read_csv(await file.read())
    identifier_type = PRODUCTS.get(c.product, PRODUCTS["other"])["identifier_type"]
    return {
        "filename": file.filename,
        "headers": headers,
        "total_rows": len(rows),
        "sample_rows": rows[:5],
        "suggested_mapping": _suggest_mapping(headers, identifier_type),
        "identifier_type": identifier_type,
    }


@router.post("/{campaign_id}/leads/upload", response_model=UploadResultOut)
async def upload_leads(campaign_id: str, file: UploadFile = File(...),
                       mapping: str = Form(...),
                       user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Store leads from a CSV using the confirmed column mapping.

    mapping: JSON object {csv_column: target}; target ∈ name | phone_number |
    bank_name | state | product_identifier | agent_field | ignore.
    Rows with no phone number are invalid; a phone already in this campaign
    (or repeated in the file) is a duplicate and is skipped — the same
    behaviour the wireframes show ("858 valid · 2 duplicates skipped").
    """
    c = _get_scoped(db, user, campaign_id)
    try:
        mapping = json.loads(mapping)
    except (TypeError, ValueError):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "mapping must be a JSON object")
    if not isinstance(mapping, dict) or "phone_number" not in mapping.values():
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "mapping must map one column to phone_number")
    bad = {t for t in mapping.values() if t not in ALL_TARGETS}
    if bad:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, f"unknown mapping targets: {sorted(bad)}")

    headers, rows = _read_csv(await file.read())
    identifier_type = PRODUCTS.get(c.product, PRODUCTS["other"])["identifier_type"]
    seen = {p for (p,) in db.query(Lead.phone_number).filter(Lead.campaign_id == c.campaign_id)}
    created = duplicates = invalid = 0
    for row in rows:
        fields = {"name": "", "phone_number": "", "bank_name": "", "state": ""}
        agent_fields = {}
        for col, target in mapping.items():
            if target == "ignore" or col not in row:
                continue
            value = row[col]
            if target in fields:
                fields[target] = value
            elif target == "product_identifier":
                if value:
                    agent_fields["product_identifier"] = value
                    agent_fields["product_identifier_type"] = identifier_type
            else:  # agent_field — keep under the normalised column name
                if value:
                    agent_fields[_norm_header(col)] = value
        phone = re.sub(r"\s+", "", fields["phone_number"])
        if not phone:
            invalid += 1
            continue
        if phone in seen:
            duplicates += 1
            continue
        seen.add(phone)
        db.add(Lead(
            tenant_id=user.tenant_id, campaign_id=c.campaign_id,
            phone_number=phone, name=fields["name"], bank_name=fields["bank_name"],
            state=fields["state"], agent_fields=agent_fields,
        ))
        created += 1
    db.commit()
    return UploadResultOut(total_rows=len(rows), created=created,
                           duplicates_skipped=duplicates, invalid_skipped=invalid)
